# -*- coding: utf-8 -*-
"""Build midterm form, slides, script, and Q&A. Writes only under midterm/."""

from __future__ import annotations

import hashlib
import shutil
from copy import deepcopy
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.patches import FancyBboxPatch, FancyArrowPatch
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn as dqn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor as DocxRGB
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt as PptPt

ROOT = Path(r"D:\study\科研拔尖项目")
MID = ROOT / "midterm"
ASSETS = MID / "assets"
NAVY = (0x1D, 0x28, 0x71)
BLUE = (0x38, 0x42, 0x82)
PALE = (0xE6, 0xE8, 0xF4)
INK = (0x1A, 0x1A, 0x1A)
GOLD = (0xC4, 0xA3, 0x5A)
WHITE = (255, 255, 255)
RED = (0xA3, 0x3B, 0x3B)
GREEN = (0x1F, 0x6B, 0x45)

FONT = "Microsoft YaHei"
FORM_FONT = "仿宋_GB2312"


def han_count(text: str) -> int:
    return sum(1 for ch in text if "\u4e00" <= ch <= "\u9fff")


def char_count(text: str) -> int:
    return sum(1 for ch in text if not ch.isspace())


def setup_mpl() -> None:
    plt.rcParams["font.sans-serif"] = [FONT, "Microsoft YaHei", "SimHei"]
    plt.rcParams["axes.unicode_minus"] = False
    plt.rcParams["font.size"] = 14


def _box(ax, x, y, w, h, text, fc, ec=NAVY, tc=WHITE, size=13, weight="bold"):
    patch = FancyBboxPatch(
        (x, y),
        w,
        h,
        boxstyle="round,pad=0.012,rounding_size=0.08",
        facecolor=tuple(c / 255 for c in fc),
        edgecolor=tuple(c / 255 for c in ec),
        linewidth=1.4,
    )
    ax.add_patch(patch)
    ax.text(
        x + w / 2,
        y + h / 2,
        text,
        ha="center",
        va="center",
        fontsize=size,
        color=tuple(c / 255 for c in tc),
        fontweight=weight,
        wrap=True,
        linespacing=1.25,
    )


def _arrow(ax, x1, y1, x2, y2):
    ax.add_patch(
        FancyArrowPatch(
            (x1, y1),
            (x2, y2),
            arrowstyle="-|>",
            mutation_scale=14,
            linewidth=1.6,
            color=tuple(c / 255 for c in NAVY),
        )
    )


def new_fig():
    fig, ax = plt.subplots(figsize=(9.69, 6.0), dpi=160)
    ax.set_xlim(0, 10)
    ax.set_ylim(0, 6.2)
    ax.axis("off")
    fig.patch.set_facecolor("white")
    return fig, ax


def save_fig(fig, name: str) -> Path:
    path = ASSETS / name
    fig.savefig(path, dpi=160)
    plt.close(fig)
    return path


def fig_order() -> None:
    fig, ax = new_fig()
    ax.text(5, 5.85, "同一道题，只改前提顺序", ha="center", va="center", fontsize=18, color=tuple(c / 255 for c in NAVY), fontweight="bold")
    _box(ax, 0.45, 4.55, 4.2, 0.85, "证明顺序", NAVY, size=16)
    _box(ax, 5.35, 4.55, 4.2, 0.85, "打乱之后", RED, ec=RED, size=16)
    for i, label in enumerate(["规则 A：前提", "规则 B：由 A 推出", "规则 C：由 B 推出"]):
        _box(ax, 0.7, 3.35 - i * 1.05, 3.7, 0.8, label, PALE, tc=NAVY, size=13, weight="normal")
    shuffled = ["规则 C：由 B 推出", "规则 A：前提", "规则 B：由 A 推出"]
    for i, label in enumerate(shuffled):
        _box(ax, 5.6, 3.35 - i * 1.05, 3.7, 0.8, label, (0xF8, 0xE8, 0xE8), ec=RED, tc=RED, size=13, weight="normal")
    ax.text(5, 0.28, "Chen et al., ICML 2024：顺序偏离时，准确率可降 30% 以上", ha="center", fontsize=12, color=tuple(c / 255 for c in INK))
    save_fig(fig, "fig_order.png")


def fig_goal() -> None:
    fig, ax = new_fig()
    ax.text(5, 5.85, "要分开的两件事", ha="center", fontsize=18, color=tuple(c / 255 for c in NAVY), fontweight="bold")
    _box(ax, 0.55, 3.15, 4.15, 2.15, "同图稳定\n换一种写法\n答案会不会翻", NAVY, size=16)
    _box(ax, 5.3, 3.15, 4.15, 2.15, "新图迁移\n换一种依赖结构\n能力还在不在", BLUE, size=16)
    _box(ax, 1.3, 0.7, 7.4, 1.55, "然后才问：训练能不能把这种结构鲁棒性提上去", PALE, tc=NAVY, size=15)
    _arrow(ax, 2.6, 3.15, 3.6, 2.25)
    _arrow(ax, 7.4, 3.15, 6.4, 2.25)
    save_fig(fig, "fig_goal.png")


def fig_2x2() -> None:
    fig, ax = new_fig()
    ax.text(5, 5.9, "二乘二：稳定还是迁移", ha="center", fontsize=18, color=tuple(c / 255 for c in NAVY), fontweight="bold")
    headers = [(3.15, "规范呈现"), (6.55, "偏移呈现")]
    for x, t in headers:
        _box(ax, x, 4.7, 2.9, 0.75, t, NAVY, size=14)
    rows = [
        (3.55, "已见图族", "同图稳定  SS", "同图翻转  SF", NAVY),
        (1.85, "未见图族", "换图差距  ΔT", "交互项", BLUE),
    ]
    for y, lab, a, b, color in rows:
        _box(ax, 0.25, y, 2.6, 1.35, lab, PALE, tc=NAVY, size=14)
        _box(ax, 3.15, y, 2.9, 1.35, a, color, size=13)
        _box(ax, 6.35, y, 3.35, 1.35, b, color, size=13)
    ax.text(5, 0.85, "重采样单位是潜在任务，不是单条文本", ha="center", fontsize=13, color=tuple(c / 255 for c in INK))
    ax.text(5, 0.35, "潜在任务 z =（依赖图, 问题, 答案）    呈现只改顺序", ha="center", fontsize=13, color=tuple(c / 255 for c in NAVY))
    save_fig(fig, "fig_2x2.png")


def fig_bgm() -> None:
    fig, ax = new_fig()
    ax.text(5, 5.9, "负担引导挖掘", ha="center", fontsize=18, color=tuple(c / 255 for c in NAVY), fontweight="bold")
    steps = [
        (0.3, "潜在任务"),
        (2.7, "多个等价\n呈现"),
        (5.1, "负担分\n+ 是否翻转"),
        (7.5, "挑出训练\n呈现"),
    ]
    for i, (x, t) in enumerate(steps):
        _box(ax, x, 3.55, 2.15, 1.45, t, NAVY if i < 3 else GREEN, ec=NAVY if i < 3 else GREEN, size=14)
        if i < 3:
            _arrow(ax, x + 2.15, 4.25, x + 2.55, 4.25)
    _box(ax, 0.3, 1.7, 9.35, 1.25, "负担只看位置：前提距离  ·  回读次数  ·  活动前沿", PALE, tc=NAVY, size=14)
    ax.text(5, 0.85, "对照：均匀随机重排（He et al., EMNLP 2025）", ha="center", fontsize=14, color=tuple(c / 255 for c in INK))
    ax.text(5, 0.35, "同一批任务、同样条数、同样 LoRA 预算", ha="center", fontsize=13, color=tuple(c / 255 for c in NAVY))
    save_fig(fig, "fig_bgm.png")


def fig_pilot() -> None:
    fig, ax = plt.subplots(figsize=(9.69, 6.0), dpi=160)
    fig.patch.set_facecolor("white")
    labels = ["已见\n规范", "已见\n偏移", "未见\n规范", "未见\n偏移"]
    vals = [0.76, 0.76, 0.76, 0.76]
    colors = [tuple(c / 255 for c in NAVY)] * 4
    bars = ax.bar(labels, vals, color=colors, width=0.62, zorder=2)
    ax.axhline(0.76, color=tuple(c / 255 for c in RED), linestyle="--", linewidth=1.6, zorder=3)
    ax.text(3.55, 0.80, "多数类基线  0.76", color=tuple(c / 255 for c in RED), fontsize=12, ha="right")
    ax.set_ylim(0, 1.05)
    ax.set_ylabel("准确率", fontsize=13)
    ax.set_title("试点四格都是 0.76，等于多数类", fontsize=18, color=tuple(c / 255 for c in NAVY), pad=12, fontweight="bold")
    for b in bars:
        ax.text(b.get_x() + b.get_width() / 2, 0.78, "0.76", ha="center", va="bottom", fontsize=12, color=tuple(c / 255 for c in NAVY))
    ax.tick_params(labelsize=12)
    ax.spines["top"].set_visible(False)
    ax.spines["right"].set_visible(False)
    fig.text(0.5, 0.045, "Qwen2.5-1.5B · 200 条 · 每格 yes 38 / no 12 · 只生成 4 个 token", ha="center", fontsize=11, color=tuple(c / 255 for c in INK))
    fig.text(0.5, 0.012, "来源：experiments/results/mvp_qwen25_1.5b/metrics.json", ha="center", fontsize=9, color="#555555")
    fig.tight_layout(rect=(0, 0.07, 1, 1))
    fig.savefig(ASSETS / "fig_pilot.png", dpi=160)
    plt.close(fig)


def fig_status() -> None:
    fig, ax = new_fig()
    ax.text(5, 5.75, "现在走到哪一步", ha="center", fontsize=18, color=tuple(c / 255 for c in NAVY), fontweight="bold")
    rows = [
        (4.35, GREEN, "已完成", "200 条试点：测量没有生效"),
        (3.15, (0xB8, 0x86, 0x0B), "进行中", "E0 健全性检查，指标未落盘"),
        (1.95, BLUE, "未开始", "训练比较、消融和边界实验"),
        (0.75, PALE, "因此", "今天不报告新的准确率", NAVY),
    ]
    for y, color, tag, desc, *rest in rows:
        tc = rest[0] if rest else WHITE
        tag_tc = NAVY if color == PALE else WHITE
        _box(ax, 0.4, y, 2.3, 0.95, tag, color, ec=color if color != PALE else NAVY, tc=tag_tc, size=15)
        _box(ax, 2.9, y, 6.6, 0.95, desc, PALE, tc=NAVY, size=15, weight="normal")
    save_fig(fig, "fig_status.png")


def fig_gantt() -> None:
    fig, ax = plt.subplots(figsize=(9.69, 6.0), dpi=160)
    fig.patch.set_facecolor("white")
    tasks = [
        ("E0 审核", 0, 7, NAVY),
        ("E1 试跑", 7, 19, BLUE),
        ("答辩（上旬）", 20, 29, (0x5C, 0x6B, 0xB0)),
        ("E1 全量", 32, 50, (0x8E, 0x99, 0xC8)),
        ("E2 / E3", 50, 68, (0xB7, 0xC0, 0xDE)),
        ("论文与代码", 68, 92, (0xD5, 0xDA, 0xEE)),
    ]
    for i, (name, a, b, color) in enumerate(tasks):
        ax.barh(i, b - a, left=a, height=0.55, color=tuple(c / 255 for c in color), zorder=2)
        ax.text(-1.5, i, name, ha="right", va="center", fontsize=12, color=tuple(c / 255 for c in NAVY))
    ax.axvline(19, color=tuple(c / 255 for c in RED), linestyle="--", linewidth=1.4, zorder=3)
    ax.text(19, 5.55, "10/31 交表", color=tuple(c / 255 for c in RED), fontsize=11, ha="center")
    ax.set_yticks([])
    ax.set_xlim(-2, 96)
    ax.set_ylim(-0.7, 6.2)
    ax.set_xticks([0, 19, 24, 45, 68, 92])
    ax.set_xticklabels(["10/12", "10/31", "11月上旬", "11月下旬", "12月", "其后"], fontsize=11)
    ax.set_title("从交表日倒排", fontsize=18, color=tuple(c / 255 for c in NAVY), fontweight="bold", pad=10)
    ax.spines["top"].set_visible(False)
    ax.spines["right"].set_visible(False)
    ax.spines["left"].set_visible(False)
    fig.tight_layout()
    fig.savefig(ASSETS / "fig_gantt.png", dpi=160)
    plt.close(fig)


def fig_risks() -> None:
    fig, ax = new_fig()
    ax.text(5, 5.8, "风险与退路", ha="center", fontsize=18, color=tuple(c / 255 for c in NAVY), fontweight="bold")
    items = [
        (4.15, "算力", "约 8GB 显存，先做 1.5B"),
        (2.55, "方法", "若不优于随机重排，就写诊断"),
        (0.95, "外推", "合成题还要对照公开基准"),
    ]
    for y, tag, desc in items:
        _box(ax, 0.4, y, 2.2, 1.2, tag, NAVY, size=16)
        _box(ax, 2.85, y, 6.7, 1.2, desc, PALE, tc=NAVY, size=15, weight="normal")
    save_fig(fig, "fig_risks.png")


# --- spoken script: non-space character count is checked later ---
SLIDES = [
    {
        "title": "封面",
        "note": "各位老师好。我是邱添，指导教师滕佳烨。下面汇报结构鲁棒性对大模型推理能力提升的探究。",
    },
    {
        "title": "问题与目标",
        "note": "做对一道题，不等于换一种等价写法还对。Chen等人在ICML 2024发现，前提顺序一偏离证明顺序，准确率可以下降百分之三十以上。我们要把同图上的呈现稳定，和换到新依赖结构之后还在不在，分开来看。",
    },
    {
        "title": "前期工作与重新定位",
        "note": "随机重排后再微调，文献里已经有了。He等人在EMNLP 2025报告，均匀随机重排在各档最好。我们的试点是Qwen2.5的1.5B、二百条，四格准确率都是0.76，等于多数类，因为当时只生成四个token。所以改做技术论文，检验负担引导能不能超过随机重排。",
    },
    {
        "title": "诊断与方法",
        "note": "测量用二乘二：已见或未见图族，乘以规范或偏移呈现。同图看会不会翻转，换图看能不能迁移。方法上按前提距离、回读次数和活动前沿来挑训练呈现，并看模型当前有没有翻转。各条件预算相同，结果高低都如实写。",
    },
    {
        "title": "实验与当前进度",
        "note": "正式的健全性检查正在运行，思维链已经打开，指标还没落盘，所以今天不报新的准确率。刚才的0.76只说明上一轮没有测到推理。训练比较还没开始。请以落盘文件为准。",
    },
    {
        "title": "风险、时间与成果",
        "note": "风险是八GB显存，以及负担引导可能赢不了随机重排；赢不了就改写成诊断论文。表格十月三十一日前提交，答辩在十一月上旬，具体日子还没定。预期是一篇中英文技术论文和开源代码，目前没有发表成果。",
    },
    {
        "title": "致谢",
        "note": "以上是中期进展，请老师批评指正。",
    },
]


def build_figures() -> None:
    setup_mpl()
    ASSETS.mkdir(parents=True, exist_ok=True)
    fig_order()
    fig_goal()
    fig_2x2()
    fig_bgm()
    fig_pilot()
    fig_status()
    fig_gantt()
    fig_risks()


def set_run_font_pptx(run, name, size, bold, color):
    run.font.name = name
    run.font.size = PptPt(size)
    run.font.bold = bold
    if color is not None:
        run.font.color.rgb = RGBColor(*color)
    rPr = run._r.get_or_add_rPr()
    for tag in ("a:latin", "a:ea", "a:cs"):
        node = rPr.find(qn(tag))
        if node is None:
            node = rPr.makeelement(qn(tag), {})
            rPr.append(node)
        node.set("typeface", name)


def set_shape_text(shape, lines, size=16, bold=False, color=NAVY, align="left", font="微软雅黑"):
    if not shape.has_text_frame:
        return
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.auto_size = None
    align_map = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}
    if isinstance(lines, str):
        lines = [lines]
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align_map[align]
        p.clear()
        run = p.add_run()
        run.text = line
        set_run_font_pptx(run, font, size, bold, color)


def duplicate_slide(prs, index):
    source = prs.slides[index]
    dest = prs.slides.add_slide(source.slide_layout)
    sp_tree = dest.shapes._spTree
    for child in list(sp_tree):
        if child.tag.endswith("}nvGrpSpPr") or child.tag.endswith("}grpSpPr"):
            continue
        sp_tree.remove(child)
    rid_map = {}
    for rel in source.part.rels.values():
        if "notesSlide" in rel.reltype:
            continue
        rid_map[rel.rId] = dest.part.relate_to(rel.target_part, rel.reltype)
    rel_ns = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
    for child in source.shapes._spTree:
        if child.tag.endswith("}nvGrpSpPr") or child.tag.endswith("}grpSpPr"):
            continue
        new_el = deepcopy(child)
        for node in new_el.iter():
            for attr_name, attr_val in list(node.attrib.items()):
                if attr_name.startswith(rel_ns) and attr_val in rid_map:
                    node.attrib[attr_name] = rid_map[attr_val]
        sp_tree.append(new_el)
    return dest


def delete_slide(prs, index):
    sld_id_lst = prs.slides._sldIdLst
    sld_id = list(sld_id_lst)[index]
    r_id = sld_id.get(qn("r:id"))
    prs.part.drop_rel(r_id)
    sld_id_lst.remove(sld_id)


def shapes_with_text(slide):
    return [sh for sh in slide.shapes if sh.has_text_frame and sh.text_frame.text.strip()]


def find_title_bars(slide):
    hits = []
    for sh in slide.shapes:
        if not sh.has_text_frame:
            continue
        if sh.top < 1_200_000 and sh.width > 3_000_000 and "请" in sh.text_frame.text:
            hits.append(sh)
        elif sh.top < 1_200_000 and sh.width > 3_000_000 and "标题" in sh.text_frame.text:
            hits.append(sh)
    return hits


def find_badge(slide):
    hits = []
    for sh in slide.shapes:
        if not sh.has_text_frame:
            continue
        if sh.left < 2_000_000 and sh.top < 1_600_000 and sh.width < 2_000_000:
            text = sh.text_frame.text.strip()
            if text and len(text) <= 4:
                hits.append(sh)
    return hits


def two_frames(slide):
    frames = []
    for sh in slide.shapes:
        if sh.width > 4_000_000 and 2_000_000 < sh.height < 4_000_000 and sh.top > 1_200_000:
            frames.append(sh)
    frames.sort(key=lambda s: s.left)
    return frames


def caption_heads(slide):
    heads = []
    for sh in slide.shapes:
        if not sh.has_text_frame:
            continue
        if 4_200_000 < sh.top < 4_900_000 and sh.height < 700_000:
            heads.append(sh)
    heads.sort(key=lambda s: s.left)
    return heads


def caption_bodies(slide):
    bodies = []
    for sh in slide.shapes:
        if not sh.has_text_frame:
            continue
        if sh.top > 4_900_000 and sh.height > 800_000:
            bodies.append(sh)
    bodies.sort(key=lambda s: s.left)
    return bodies


def place_picture(slide, frame, image):
    pad = Emu(70_000)
    slide.shapes.add_picture(str(image), frame.left + pad, frame.top + pad, frame.width - 2 * pad, frame.height - 2 * pad)
    if frame.has_text_frame:
        set_shape_text(frame, "", size=12, color=WHITE)


def style_two_figure(slide, badge, title, images, heads, bodies):
    for sh in find_badge(slide):
        set_shape_text(sh, badge, size=20, bold=True, color=NAVY, align="center")
    for sh in find_title_bars(slide):
        set_shape_text(sh, title, size=22, bold=True, color=WHITE, align="center")
    frames = two_frames(slide)
    if len(frames) < 2:
        raise RuntimeError(f"expected 2 frames, got {len(frames)} on {title}")
    for frame, image in zip(frames, images):
        place_picture(slide, frame, ASSETS / image)
    hs = caption_heads(slide)
    bs = caption_bodies(slide)
    for sh, text in zip(hs, heads):
        set_shape_text(sh, text, size=15, bold=True, color=NAVY, align="left")
    for sh, text in zip(bs, bodies):
        set_shape_text(sh, text, size=14, bold=False, color=INK, align="left")


def card_titles(slide):
    hits = [sh for sh in slide.shapes if sh.has_text_frame and 2_400_000 < sh.top < 2_950_000 and sh.width < 3_500_000]
    hits.sort(key=lambda s: s.left)
    return hits


def card_bodies(slide):
    hits = [sh for sh in slide.shapes if sh.has_text_frame and 2_950_000 < sh.top < 3_600_000 and sh.width < 4_000_000]
    hits.sort(key=lambda s: (s.left, s.top))
    # template repeats a wide body under each card; keep the three whose left matches titles
    titles = card_titles(slide)
    chosen = []
    for t in titles:
        near = [b for b in hits if abs(b.left - t.left) < 400_000]
        if near:
            chosen.append(sorted(near, key=lambda s: s.top)[0])
    return chosen


def fill_pills(slide, date_text, speaker, college):
    pills = []
    for sh in slide.shapes:
        if not sh.has_text_frame:
            continue
        raw = sh.text_frame.text
        if "汇报人" in raw or "学院" in raw or "月" in raw or "xx" in raw.lower() or "xxx" in raw.lower():
            if sh.top > 4_500_000 and sh.height < 800_000:
                pills.append(sh)
    pills.sort(key=lambda s: s.left)
    texts = [date_text, speaker, college]
    if len(pills) < 3:
        raise RuntimeError(f"expected 3 pills, found {len(pills)}: {[p.text_frame.text for p in pills]}")
    sizes = [13, 15, 15]
    for sh, text, size in zip(pills, texts, sizes):
        set_shape_text(sh, text, size=size, bold=True, color=NAVY, align="center")


def big_title(slide):
    cands = []
    for sh in slide.shapes:
        if not sh.has_text_frame:
            continue
        if sh.width > 5_000_000 and 2_000_000 < sh.top < 4_500_000:
            cands.append(sh)
    if not cands:
        raise RuntimeError("cover title box not found")
    return max(cands, key=lambda s: s.width)


def add_subtitle(slide, text):
    box = slide.shapes.add_textbox(Emu(2_200_000), Emu(4_180_000), Emu(7_800_000), Emu(420_000))
    set_shape_text(box, text, size=16, bold=False, color=WHITE, align="center")
    box.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER


def set_notes(slide, text):
    notes = slide.notes_slide
    tf = notes.notes_text_frame
    tf.text = text
    for p in tf.paragraphs:
        for run in p.runs:
            set_run_font_pptx(run, "微软雅黑", 14, False, INK)


def build_ppt() -> Path:
    tpl = next(p for p in ROOT.glob("*.pptx") if p.parent == ROOT)
    out = MID / "中期答辩.pptx"
    shutil.copyfile(tpl, out)
    prs = Presentation(str(out))
    # indices in the original template
    cover = duplicate_slide(prs, 11)
    problem = duplicate_slide(prs, 6)
    prior = duplicate_slide(prs, 7)
    method = duplicate_slide(prs, 6)
    progress = duplicate_slide(prs, 6)
    plan = duplicate_slide(prs, 6)
    thanks = duplicate_slide(prs, 11)
    n_orig = 12
    for _ in range(n_orig):
        delete_slide(prs, 0)

    # cover
    set_shape_text(
        big_title(cover),
        ["结构鲁棒性对大模型", "推理能力提升的探究"],
        size=32,
        bold=True,
        color=WHITE,
        align="center",
    )
    add_subtitle(cover, "拔尖型科研训练项目  ·  中期答辩")
    fill_pills(cover, "2026年11月上旬", "汇报人：邱添", "统计与管理学院")

    style_two_figure(
        problem,
        "一",
        "问题与目标",
        ["fig_order.png", "fig_goal.png"],
        ["顺序一变", "目标"],
        ["准确率可降 30% 以上", "分开同图稳定与新图迁移"],
    )
    # prior: three cards
    for sh in find_badge(prior):
        set_shape_text(sh, "二", size=20, bold=True, color=NAVY, align="center")
    for sh in find_title_bars(prior):
        set_shape_text(sh, "前期工作与重新定位", size=20, bold=True, color=WHITE, align="center")
    titles = ["文献", "试点", "转向"]
    bodies = ["重排微调已有工作", "准确率等于多数类", "改检验负担引导"]
    ct = card_titles(prior)
    cb = card_bodies(prior)
    if len(ct) < 3 or len(cb) < 3:
        raise RuntimeError(f"cards titles={len(ct)} bodies={len(cb)}")
    for sh, text in zip(ct, titles):
        set_shape_text(sh, text, size=20, bold=True, color=WHITE, align="center")
    for sh, text in zip(cb, bodies):
        set_shape_text(sh, text, size=16, bold=False, color=WHITE, align="center")

    style_two_figure(
        method,
        "三",
        "诊断与方法",
        ["fig_2x2.png", "fig_bgm.png"],
        ["怎么测", "怎么训"],
        ["同图翻转，还是换图失败", "等预算对照随机重排"],
    )
    style_two_figure(
        progress,
        "四",
        "实验与当前进度",
        ["fig_pilot.png", "fig_status.png"],
        ["已有试点", "正式结果"],
        ["0.76 不是结构鲁棒", "E0 进行中，训练未开始"],
    )
    style_two_figure(
        plan,
        "五",
        "风险、时间与成果",
        ["fig_gantt.png", "fig_risks.png"],
        ["时间", "退路"],
        ["10月31日交表，11月上旬答辩", "不赢就写诊断，不改数据"],
    )

    set_shape_text(big_title(thanks), ["感谢聆听", "请批评指正"], size=36, bold=True, color=WHITE, align="center")
    add_subtitle(thanks, "指导教师：滕佳烨")
    fill_pills(thanks, "2026年11月上旬", "汇报人：邱添", "统计与管理学院")

    slides = [cover, problem, prior, method, progress, plan, thanks]
    for slide, meta in zip(slides, SLIDES):
        set_notes(slide, meta["note"])

    leftovers = []
    for i, slide in enumerate(prs.slides):
        for sh in slide.shapes:
            if not sh.has_text_frame:
                continue
            raw = sh.text_frame.text
            if any(k in raw for k in ("请添加", "请输入", "xxx", "在此处", "在此添加")):
                leftovers.append((i, raw[:40]))
    if leftovers:
        raise RuntimeError(f"placeholder text remains: {leftovers}")

    prs.core_properties.title = "结构鲁棒性对大模型推理能力提升的探究"
    prs.core_properties.author = "邱添"
    prs.core_properties.subject = "拔尖型科研训练项目中期答辩"
    prs.save(str(out))
    return out


def tr_cells(table, row_idx):
    from docx.table import _Cell

    tr = table.rows[row_idx]._tr
    return [_Cell(tc, table) for tc in tr.findall(dqn("w:tc"))]


def shade_cell(cell, fill="FFFF00"):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = tcPr.find(dqn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tcPr.append(shd)
    shd.set(dqn("w:val"), "clear")
    shd.set(dqn("w:color"), "auto")
    shd.set(dqn("w:fill"), fill)


def write_run(paragraph, text, size=12, bold=False):
    paragraph.clear()
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.size = Pt(size)
    run.font.name = FORM_FONT
    run.font.color.rgb = DocxRGB(0, 0, 0)
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(dqn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rFonts.set(dqn(attr), FORM_FONT)
    pf = paragraph.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(2)
    pf.line_spacing = 1.15
    pf.alignment = WD_ALIGN_PARAGRAPH.LEFT


def trim_trailing_empty(cell):
    while len(cell.paragraphs) > 1 and not cell.paragraphs[-1].text.strip():
        element = cell.paragraphs[-1]._element
        element.getparent().remove(element)


def fill_lines(cell, lines, keep_first=False):
    start = 1 if keep_first else 0
    while len(cell.paragraphs) < start + len(lines):
        cell.add_paragraph("")
    paras = cell.paragraphs
    for i, line in enumerate(lines):
        write_run(paras[start + i], line, size=12, bold=False)
    # Drop the template's spare blank lines so filled sections do not spill an empty page.
    trim_trailing_empty(cell)


PROGRESS = [
    "本项目研究大模型推理的结构鲁棒性：同一道题只改变前提的合法写法时，答案是否还稳定；这种稳定是否还能迁移到没见过的依赖结构。Chen 等人（Premise Order Matters，ICML 2024）表明，前提顺序偏离证明顺序时，准确率可以下降 30% 以上。题目得分高，不能直接当成结构上稳健的推理。",
    "项目自 2026 年 5 月开展，指导教师为滕佳烨。已完成三部分工作。",
    "第一，文献调研。诊断方面最接近的工作有 Premise Order Matters（Chen et al., ICML 2024）、DyVal（Zhu et al., ICLR 2024）、GSM-∞（Zhou et al., ICML 2025）和 iGSM（Ye et al., ICLR 2025）。用随机重排做监督微调，已有 He 等人（EMNLP 2025）和 MEND（Findings of ACL 2025）。He 等人报告均匀随机重排在各难度档上都最好。据这次调研，把“同图稳定”和“新图迁移”分开、再用可解释的依赖负担挑选训练呈现，这一组合尚未见到。",
    "第二，形式化。潜在任务是依赖图、问题和求解器给出的答案；呈现只改变句子顺序，或加入不改变答案的干扰前提。同图用稳定率和翻转率，换图用迁移差距，并做“已见或未见图族 × 规范或偏移呈现”的二乘二分解。统计按潜在任务配对，不按单条文本。",
    "第三，一次真实试点。模型为 Qwen2.5-1.5B-Instruct，200 条、每格 50 条，贪心解码且最多生成 4 个 token。四格准确率都是 0.76，稳定率 1.00，翻转率 0.00，换图差距 0.00。每格标签为 38 个 yes、12 个 no，模型几乎全答 yes。因此 0.76 只等于多数类基线，不能作为结构鲁棒性的证据。来源：experiments/results/mvp_qwen25_1.5b/metrics.json，以及 experiments/mvp_pilot_report.md、run_metadata.json。",
    "据此把论文从“做基准”调整为技术论文。诊断要回答：现有的顺序不变性训练，改善的是同图稳定，还是也能迁到新的依赖结构。方法要回答：按依赖负担挑选训练呈现，是否优于均匀随机重排。负担只由相关前提的文本距离、回读次数和活动依赖前沿计算。",
    "正式健全性检查正在运行：打开思维链，每格不少于 200 条，比较证明顺序、随机顺序和逆序。该轮的 metrics.json 与 rows.jsonl 尚未落盘，本表不报告其准确率，也不用更小样本代替。求解器核对的数据生成、负担特征和配对重抽样代码已经写好。训练尚未开始。",
]

OUTCOMES = [
    "截至 2026 年 10 月 12 日，没有发表论文，没有专利、奖励、专著或媒体报道。",
    "阶段性工作是文献综述、问题形式化，以及一次说明测量尚未生效的试点。正在运行的检查不计入成果。",
]

PROBLEMS = [
    "1. 算力为一台约 8GB 显存的笔记本 GPU（试点记录中的 RTX 5060 Laptop）。主模型暂为 Qwen2.5-1.5B-Instruct。结论不能直接推广到更大模型。",
    "2. 第一次试点没有测到顺序效应。生成长度只有 4 个 token，标签也不平衡。纠正后的健全性检查仍在运行，正式数字未出。",
    "3. 负担引导挖掘可能并不优于均匀随机重排。He 等人（EMNLP 2025）的分档结果提示这一点。若不成立，改为以诊断为主的分析论文，不改写数据。",
    "4. 合成规则题不能单独代表真实数学题。R-GSM 没有找到公开文件。后续用 GSM8K 检查通用能力是否下降；数学顺序测试会标明是自建的前提重排。",
    "5. 本表没有经费栏目。经费数额未向填写人提供，因此任何位置都不填写经费，避免编造。",
]

CHANGES = [
    "指导教师未变更，仍为滕佳烨。项目组成员一栏空白，待补充，不能据此声称成员已经变更或没有变更。",
    "研究方案在组内作了调整，是否已向学院办理书面备案，填写时尚未确认，请负责人核定。调整内容是：由以评测基准为主，改为技术论文；训练比较收束为严格等预算下，用负担引导挖掘对照均匀随机重排，并用二乘二分开同图稳定与新图迁移。",
    "调整原因有两条。一是随机重排微调已有公开工作，原先更宽的方案缺少差异。二是 200 条试点没有有效信号，测量方式必须先改（打开思维链、平衡 yes/no），再谈训练。",
]

PLAN = [
    "时间按中期表格截止日期 2026 年 10 月 31 日倒排。答辩安排在 11 月上旬，具体日期未定。以下是工作计划，不是已经完成的结果。",
    "1. 10 月 12 日至 18 日：完成正在运行的健全性检查并审核。至少有一档规范呈现的准确率落在 0.55 到 0.90，且翻转率不低于 10%、置信区间不含 0，才进入训练。未达标只调整难度或更换模型，最多三轮。",
    "2. 10 月 19 日至 31 日：若检查通过，做小规模等预算试跑。条件包括不训练、规范顺序微调、均匀随机重排、随机重排加一致性项、负担引导挖掘。10 月 31 日前提交本表。",
    "3. 11 月上旬：中期答辩。自述约 3 分钟，问答约 5 分钟。",
    "4. 11 月中下旬：在约 8GB 显存允许的规模上做训练全量。种子数以实际跑完的为准，不足时如实写明。",
    "5. 12 月：做消融，分别看负担分数、错误引导和一致性项。算力允许时，在算术图上做边界实验。",
    "6. 其后：撰写中文和英文技术论文，并整理开源代码。投稿场合尚未确定。预期成果是论文和代码，不是已经发表的文章。",
]


def build_form() -> dict:
    src_docx = Path(r"C:\Users\Qiutian\AppData\Local\Temp\midterm_form.docx")
    if not src_docx.exists():
        raise FileNotFoundError(src_docx)
    out_docx = MID / "附件二 上海财经大学统计与管理学院拔尖型科研训练项目中期检查报告表_已填.docx"
    shutil.copyfile(src_docx, out_docx)
    doc = Document(str(out_docx))
    table = doc.tables[0]

    number = tr_cells(table, 0)[1]
    report_time = tr_cells(table, 0)[3]
    name = tr_cells(table, 1)[1]
    person = tr_cells(table, 2)[1]
    members = tr_cells(table, 2)[3]
    advisor = tr_cells(table, 3)[1]
    progress = tr_cells(table, 4)[1]
    outcomes = tr_cells(table, 5)[1]
    problems = tr_cells(table, 6)[1]
    changes = tr_cells(table, 7)[1]
    plan = tr_cells(table, 8)[1]

    shade_cell(number)
    shade_cell(report_time)
    shade_cell(members)

    fill_lines(name, ["结构鲁棒性对大模型推理能力提升的探究"])
    fill_lines(person, ["邱添", "2024111054"])
    fill_lines(advisor, ["滕佳烨"])
    fill_lines(progress, PROGRESS)
    fill_lines(outcomes, OUTCOMES, keep_first=True)
    fill_lines(problems, PROBLEMS)
    fill_lines(changes, CHANGES, keep_first=True)
    fill_lines(plan, PLAN)

    doc.save(str(out_docx))

    counts = {
        "项目名称": "结构鲁棒性对大模型推理能力提升的探究",
        "负责人/联系方式": "邱添，学号 2024111054",
        "指导老师": "滕佳烨",
        "一、项目进展情况汇报": "\n".join(PROGRESS),
        "二、项目取得的成果": "\n".join(OUTCOMES),
        "三、存在的问题和困难": "\n".join(PROBLEMS),
        "四、项目变更情况": "\n".join(CHANGES),
        "五、后期项目建设安排": "\n".join(PLAN),
    }
    return {"docx": out_docx, "counts": counts}


def write_markdown() -> None:
    lines = ["# 中期答辩讲稿", "", "自述控制在 3 分钟。下面每页是备注里的口播稿，按页顺序念。", ""]
    total = ""
    for i, meta in enumerate(SLIDES, start=1):
        lines.append(f"## 第 {i} 页  {meta['title']}")
        lines.append("")
        lines.append(meta["note"])
        lines.append("")
        total += meta["note"]
    lines.append("## 字数")
    lines.append("")
    lines.append(f"口播稿合计：汉字 {han_count(total)}，不含空白字符 {char_count(total)}。")
    lines.append("")
    lines.append("## 数字从哪来")
    lines.append("")
    lines.append("- 四格准确率 0.76、稳定率 1.00、翻转率 0.00、换图差距 0.00、每格 yes 38 / no 12：`experiments/results/mvp_qwen25_1.5b/metrics.json`。")
    lines.append("- 200 条、每格 50 条、最多生成 4 个 token、模型 Qwen2.5-1.5B-Instruct：`experiments/mvp_pilot_report.md` 与 `experiments/results/mvp_qwen25_1.5b/run_metadata.json`。")
    lines.append("- 前提顺序可令准确率下降 30% 以上：docs/01 已核实的 Chen et al., ICML 2024，不是本项目的实验结果。")
    lines.append("- 均匀随机重排在各档最好：docs/01 已核实的 He et al., EMNLP 2025，不是本项目的实验结果。")
    lines.append("- E0 只说“正在运行”。`build/e0/round1/` 在制表时没有 metrics.json 或 rows.jsonl。smoke 与 micro 的数字没有写入讲稿。")
    lines.append("")
    (MID / "讲稿.md").write_text("\n".join(lines), encoding="utf-8")

    qa = """# 评委可能提问及短答

问答大约 5 分钟。每条先给一句结论，再补一句依据。不要把没跑完的实验说成结果。

## 1. 和 Premise Order Matters 有什么区别？

他们证明前提顺序会改变准确率，训练留给以后做。我们做的是训练干预，而且要把同图翻转和新图失败分开测。

## 2. 和 He 等人 2025 年那篇有什么区别？

他们用均匀随机重排做微调，并发现随机在各档最好。我们把这句话当成要检验的假设，用负担分数选呈现，预算和随机重排对齐。

## 3. 和 MEND 呢？

MEND 也做顺序和冗余上的增强。据我们的调研，他们没有在同一潜在任务上估计条件翻转，也没有在生成时把训练图族和测试图族隔开。

## 4. 只用 1.5B，结论能推广吗？

不能直接推广。显存大约 8GB，所以先做 Qwen2.5-1.5B。模型规模会写进限制，不用小模型的结果代替大模型。

## 5. 如果方法无效怎么办？

退路已经定了：论文改成以诊断为主，报告随机重排改善的是同图稳定还是新图迁移。不改数据，也不把阴性结果写成提升。

## 6. 0.76 是不是说明已经很稳？

不是。四格都是 0.76，每格 38 个 yes、12 个 no，模型几乎全答 yes，而且只生成了 4 个 token。稳定率 1 只说明输出没有翻，不说明会推理。

## 7. 为什么现在还没有 E0 的数字？

这一轮还在跑，metrics.json 还没写出来。更小的试跑样本不拿来当正式结果。

## 8. 负担分数会不会只是在挑更深的题？

三个特征只看相关规则在文本里的位置：距离、回读逆序、活动前沿。深度用来分层，不放进分数。

## 9. 什么叫等预算？

同一批潜在任务，每个任务同样两条训练序列，LoRA 的秩、目标层和步数相同。差别只在选哪些呈现、加不加一致性项。

## 10. 合成题能说明真实推理吗？

不能单独说明。后面会看 GSM8K 会不会掉点。R-GSM 没有找到公开文件，数学顺序测试会做成自建的前提重排，并在文中说明。

## 11. 有没有论文，经费是多少？

没有发表、专利或奖励。本表没有经费栏，经费数额也没有提供，所以没填。

## 12. 和 DyVal、iGSM、GSM-∞ 的区别？

它们用图控制难度。它们没有把同一张图的不同呈现，做成和换图正交的因子，也没有做我们这种等预算的训练比较。

## 13. 五月到现在为什么还在重做？

文献、定义和那次试点都还在。重做是因为试点证明测量无效，而且最近的文献已经覆盖了“随便重排再微调”。中期报告的就是这次修正。

## 14. 成员和方案变更要不要备案？

指导教师没有变，仍是滕佳烨。成员名单和联系电话表里还空着。方案调整目前是组内的，要不要向学院书面备案，需要负责人确认。
"""
    (MID / "评委问答.md").write_text(qa, encoding="utf-8")


def main() -> None:
    MID.mkdir(parents=True, exist_ok=True)
    original = next(p for p in ROOT.glob("*.doc") if p.parent == ROOT and "_已填" not in p.name)
    template = next(p for p in ROOT.glob("*.pptx") if p.parent == ROOT)
    digest_before = {
        original.name: hashlib.sha256(original.read_bytes()).hexdigest(),
        template.name: hashlib.sha256(template.read_bytes()).hexdigest(),
    }
    build_figures()
    form = build_form()
    ppt = build_ppt()
    write_markdown()
    spoken = "".join(s["note"] for s in SLIDES)
    report = ["# build log", ""]
    report.append(f"口播汉字 {han_count(spoken)}  不含空白 {char_count(spoken)}")
    report.append("")
    report.append("表中未写明字数上限。各栏为本次填写正文（不含原表提示语）：")
    for key, val in form["counts"].items():
        report.append(f"- {key}：汉字 {han_count(val)}，不含空白 {char_count(val)}")
    report.append("")
    report.append(f"pptx {ppt}")
    report.append(f"docx {form['docx']}")
    digest_after = {
        original.name: hashlib.sha256(original.read_bytes()).hexdigest(),
        template.name: hashlib.sha256(template.read_bytes()).hexdigest(),
    }
    report.append(f"original unchanged {digest_before == digest_after}")
    (ASSETS / "build_log.txt").write_text("\n".join(report), encoding="utf-8")
    print("\n".join(report))


if __name__ == "__main__":
    main()
